import csv
import io
import os
import secrets
import smtplib
from email.message import EmailMessage
import sqlite3
import re
import hashlib
import json
import urllib.request
import urllib.error
from datetime import datetime
from pathlib import Path
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash

from flask import Flask, jsonify, request, Response, g
from flask_cors import CORS
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_DB = BASE_DIR.parent / "database" / "lba_rankings.db"
DB_PATH = Path(os.environ.get("LBA_DB_PATH", DEFAULT_DB))
LOGO_PATH = BASE_DIR.parent / "frontend" / "src" / "assets" / "lba-logo.png"
ADMIN_USERNAME = os.environ.get("LBA_ADMIN_USERNAME", "admin")
ADMIN_PASSWORD = os.environ.get("LBA_ADMIN_PASSWORD")
if not ADMIN_PASSWORD:
    raise RuntimeError("LBA_ADMIN_PASSWORD must be set in the environment before starting the backend.")
AUTH_TOKEN_DAYS = int(os.environ.get("LBA_AUTH_TOKEN_DAYS", "30"))
SMTP_HOST = os.environ.get("LBA_SMTP_HOST")
SMTP_PORT = int(os.environ.get("LBA_SMTP_PORT", "587"))
SMTP_USERNAME = os.environ.get("LBA_SMTP_USERNAME")
SMTP_PASSWORD = os.environ.get("LBA_SMTP_PASSWORD")
SMTP_FROM = os.environ.get("LBA_SMTP_FROM") or SMTP_USERNAME
PUBLIC_APP_URL = os.environ.get("LBA_PUBLIC_APP_URL", "")
RESEND_API_KEY = os.environ.get("RESEND_API_KEY")
EMAIL_FROM = os.environ.get("LBA_EMAIL_FROM") or SMTP_FROM or "onboarding@resend.dev"


def create_app():
    app = Flask(__name__)
    CORS(app)
    ensure_database()

    @app.route("/api/health")
    def health():
        return jsonify({"ok": True, "database": str(DB_PATH), "time": datetime.utcnow().isoformat() + "Z"})

    @app.route("/api/auth/login", methods=["POST"])
    def login():
        data = request.get_json(silent=True) or {}
        identifier = (data.get("username") or data.get("email") or "").strip()
        password = data.get("password") or ""
        if not identifier or not password:
            return jsonify({"error": "Username/email and password are required"}), 400

        with db() as con:
            user = con.execute("""
                SELECT * FROM users
                WHERE LOWER(username)=LOWER(?) OR LOWER(email)=LOWER(?)
                LIMIT 1
            """, (identifier, identifier)).fetchone()

            if not user or not user["is_active"] or not check_password_hash(user["password_hash"], password):
                return jsonify({"error": "Invalid username/email or password"}), 401
            if not user["email_verified"]:
                return jsonify({"error": "Please verify your email address before signing in."}), 403

            token = secrets.token_urlsafe(48)
            token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
            now = datetime.utcnow()
            expires = now.timestamp() + (AUTH_TOKEN_DAYS * 86400)
            expires_at = datetime.utcfromtimestamp(expires).replace(microsecond=0).isoformat() + "Z"
            now_text = now.replace(microsecond=0).isoformat() + "Z"

            con.execute("""
                INSERT INTO auth_tokens(user_id, token_hash, created_at, expires_at)
                VALUES(?,?,?,?)
            """, (user["id"], token_hash, now_text, expires_at))
            con.execute("UPDATE users SET last_login=?, updated_at=? WHERE id=?", (now_text, now_text, user["id"]))
            con.commit()

        return jsonify({
            "token": token,
            "user": public_user(user),
            "expires_at": expires_at
        })


    @app.route("/api/auth/register", methods=["POST"])
    def register():
        data = request.get_json(silent=True) or {}
        full_name = (data.get("full_name") or "").strip()
        username = (data.get("username") or "").strip()
        email = (data.get("email") or "").strip().lower()
        password = data.get("password") or ""

        if not full_name or not username or not email or len(password) < 8:
            return jsonify({"error": "Full name, username, email and a password of at least 8 characters are required"}), 400

        now = now_iso()
        with db() as con:
            exists = con.execute("SELECT id FROM users WHERE LOWER(username)=LOWER(?) OR LOWER(email)=LOWER(?)", (username, email)).fetchone()
            if exists:
                return jsonify({"error": "Username or email is already registered"}), 409

            con.execute("""
                INSERT INTO users(username,email,full_name,password_hash,role,is_active,email_verified,created_at,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?)
            """, (username, email, full_name, generate_password_hash(password), "Viewer", 1, 0, now, now))
            user_id = con.execute("SELECT last_insert_rowid()").fetchone()[0]
            raw_token = secrets.token_urlsafe(32)
            token_hash = hashlib.sha256(raw_token.encode()).hexdigest()
            expires = datetime.utcfromtimestamp(datetime.utcnow().timestamp() + 86400).replace(microsecond=0).isoformat() + "Z"
            con.execute("INSERT INTO email_tokens(user_id,token_hash,token_type,expires_at,created_at) VALUES(?,?,?,?,?)",
                        (user_id, token_hash, "verify", expires, now))
            con.commit()

        try:
            send_email(
                email,
                "Verify your LBA Admin account",
                f"Hello {full_name},\n\nYour LBA Admin account has been created. "
                f"Use this verification code/link token to verify your email:\n\n{raw_token}\n\n"
                f"This token expires in 24 hours.\n"
            )
        except Exception:
            # Do not leave an unusable unverified account behind when email delivery fails.
            with db() as con:
                con.execute("DELETE FROM email_tokens WHERE user_id=? AND token_type='verify'", (user_id,))
                con.execute("DELETE FROM users WHERE id=? AND email_verified=0", (user_id,))
                con.commit()
            return jsonify({
                "error": "Account could not be created because the verification email could not be sent. Please try again later."
            }), 503

        return jsonify({"message": "Registration successful. Check your email to verify your account."}), 201


    @app.route("/api/auth/resend-verification", methods=["POST"])
    def resend_verification():
        data = request.get_json(silent=True) or {}
        identifier = (data.get("email") or data.get("username") or "").strip().lower()
        if not identifier:
            return jsonify({"error": "Username or email is required"}), 400

        with db() as con:
            user = con.execute("""
                SELECT * FROM users
                WHERE LOWER(username)=LOWER(?) OR LOWER(email)=LOWER(?)
                LIMIT 1
            """, (identifier, identifier)).fetchone()

            if not user:
                return jsonify({"error": "If that account exists and is not verified, a new verification token has been sent."})

            if user["email_verified"]:
                return jsonify({"message": "This email is already verified. You can sign in now."})

            raw_token = secrets.token_urlsafe(32)
            token_hash = hashlib.sha256(raw_token.encode()).hexdigest()
            now = now_iso()
            expires = datetime.utcfromtimestamp(
                datetime.utcnow().timestamp() + 86400
            ).replace(microsecond=0).isoformat() + "Z"

            con.execute(
                "DELETE FROM email_tokens WHERE user_id=? AND token_type='verify'",
                (user["id"],)
            )
            con.execute(
                "INSERT INTO email_tokens(user_id,token_hash,token_type,expires_at,created_at) VALUES(?,?,?,?,?)",
                (user["id"], token_hash, "verify", expires, now)
            )
            con.commit()

        try:
            send_email(
                user["email"],
                "Your new LBA email verification token",
                f"Hello {user['full_name']},\\n\\n"
                f"Here is your new LBA Admin email verification token:\\n\\n"
                f"{raw_token}\\n\\n"
                f"This token expires in 24 hours. Any previous verification token is no longer valid.\\n"
            )
        except Exception:
            with db() as con:
                con.execute("DELETE FROM email_tokens WHERE token_hash=?", (token_hash,))
                con.commit()
            return jsonify({
                "error": "The verification email could not be sent right now. Please try again later."
            }), 503

        return jsonify({
            "message": "A new verification token has been sent to your registered email address."
        })


    @app.route("/api/auth/verify-email", methods=["POST"])
    def verify_email():
        data = request.get_json(silent=True) or {}
        raw_token = (data.get("token") or "").strip()
        if not raw_token:
            return jsonify({"error": "Verification token is required"}), 400
        token_hash = hashlib.sha256(raw_token.encode()).hexdigest()
        with db() as con:
            row = con.execute("""
                SELECT user_id FROM email_tokens
                WHERE token_hash=? AND token_type='verify' AND expires_at > ?
            """, (token_hash, now_iso())).fetchone()
            if not row:
                return jsonify({"error": "Invalid or expired verification token"}), 400
            con.execute("UPDATE users SET email_verified=1, updated_at=? WHERE id=?", (now_iso(), row["user_id"]))
            con.execute("DELETE FROM email_tokens WHERE token_hash=?", (token_hash,))
            con.commit()
        return jsonify({"message": "Email verified. You can now sign in."})


    @app.route("/api/auth/forgot-password", methods=["POST"])
    def forgot_password():
        data = request.get_json(silent=True) or {}
        email = (data.get("email") or "").strip().lower()
        if not email:
            return jsonify({"error": "Email is required"}), 400
        with db() as con:
            user = con.execute("SELECT * FROM users WHERE LOWER(email)=LOWER(?) AND is_active=1", (email,)).fetchone()
            if user:
                raw_token = secrets.token_urlsafe(32)
                token_hash = hashlib.sha256(raw_token.encode()).hexdigest()
                now = now_iso()
                expires = datetime.utcfromtimestamp(datetime.utcnow().timestamp() + 3600).replace(microsecond=0).isoformat() + "Z"
                con.execute("DELETE FROM email_tokens WHERE user_id=? AND token_type='reset'", (user["id"],))
                con.execute("INSERT INTO email_tokens(user_id,token_hash,token_type,expires_at,created_at) VALUES(?,?,?,?,?)",
                            (user["id"], token_hash, "reset", expires, now))
                con.commit()
                try:
                    send_email(
                        user["email"],
                        "Reset your LBA Admin password",
                        f"Hello {user['full_name']},\n\nA password reset was requested for your LBA Admin account.\n\n"
                        f"Use this reset token in the app:\n\n{raw_token}\n\n"
                        f"The token expires in 1 hour. If you did not request this, ignore this email.\n"
                    )
                except Exception:
                    # Remove the unusable reset token so a failed delivery cannot leave a misleading reset state.
                    con.execute("DELETE FROM email_tokens WHERE token_hash=?", (token_hash,))
                    con.commit()
                    return jsonify({
                        "error": "The password reset email could not be sent right now. Please try again later."
                    }), 503
        return jsonify({"message": "If that email is registered, a password reset message has been sent."})


    @app.route("/api/auth/reset-password", methods=["POST"])
    def reset_password():
        data = request.get_json(silent=True) or {}
        raw_token = (data.get("token") or "").strip()
        password = data.get("password") or ""
        if not raw_token or len(password) < 8:
            return jsonify({"error": "Reset token and a password of at least 8 characters are required"}), 400
        token_hash = hashlib.sha256(raw_token.encode()).hexdigest()
        with db() as con:
            row = con.execute("""
                SELECT user_id FROM email_tokens
                WHERE token_hash=? AND token_type='reset' AND expires_at > ?
            """, (token_hash, now_iso())).fetchone()
            if not row:
                return jsonify({"error": "Invalid or expired reset token"}), 400
            now = now_iso()
            con.execute("UPDATE users SET password_hash=?, email_verified=1, updated_at=? WHERE id=?",
                        (generate_password_hash(password), now, row["user_id"]))
            con.execute("DELETE FROM email_tokens WHERE token_hash=?", (token_hash,))
            con.execute("DELETE FROM auth_tokens WHERE user_id=?", (row["user_id"],))
            con.commit()
        return jsonify({"message": "Password reset successfully. Please sign in again."})


    @app.route("/api/auth/me")
    @require_auth
    def auth_me():
        return jsonify({"user": public_user(g.current_user)})


    @app.route("/api/auth/logout", methods=["POST"])
    @require_auth
    def auth_logout():
        token = bearer_token()
        token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
        with db() as con:
            con.execute("DELETE FROM auth_tokens WHERE token_hash=?", (token_hash,))
            con.commit()
        return jsonify({"logged_out": True})

    @app.route("/api/dashboard")
    @require_auth
    def dashboard():
        with db() as con:
            total = con.execute("SELECT COUNT(*) FROM players").fetchone()[0]
            active = con.execute("SELECT COUNT(*) FROM players WHERE status='Active'").fetchone()[0]
            categories = con.execute("SELECT COUNT(DISTINCT category_code) FROM players WHERE category_code IS NOT NULL AND category_code!=''").fetchone()[0]
            top = con.execute("SELECT id, full_name, category_code, club, rank_position, total_points FROM players ORDER BY COALESCE(rank_position, 999999), total_points DESC, full_name LIMIT 6").fetchall()
            categories_breakdown = con.execute("""
                SELECT category_code, event_type, age_group, COUNT(*) AS player_count
                FROM players WHERE category_code IS NOT NULL AND category_code!=''
                GROUP BY category_code, event_type, age_group
                ORDER BY player_count DESC, category_code
            """).fetchall()
            recent = con.execute("SELECT action, details, created_at FROM audit_logs ORDER BY id DESC LIMIT 6").fetchall()
            draws = con.execute("SELECT COUNT(*) FROM draws").fetchone()[0]
        return jsonify({
            "totalPlayers": total,
            "activePlayers": active,
            "categories": categories,
            "draws": draws,
            "topPlayer": dict(top[0]) if top else None,
            "topPlayers": [dict(row) for row in top],
            "categoryBreakdown": [dict(row) for row in categories_breakdown],
            "recentActivity": [dict(row) for row in recent],
        })

    @app.route("/api/categories")
    @require_auth
    def categories():
        with db() as con:
            rows = con.execute("""
                SELECT category_code, event_type, gender, age_group, COUNT(*) AS player_count
                FROM players
                GROUP BY category_code, event_type, gender, age_group
                ORDER BY event_type, age_group
            """).fetchall()
        return jsonify([dict(row) for row in rows])

    @app.route("/api/players", methods=["GET"])
    @require_auth
    def list_players():
        category = request.args.get("category", "").strip()
        q = request.args.get("q", "").strip()
        status = request.args.get("status", "").strip()
        sql = "SELECT * FROM players WHERE 1=1"
        args = []
        if category and category != "All":
            sql += " AND category_code = ?"
            args.append(category)
        if status and status != "All":
            sql += " AND status = ?"
            args.append(status)
        if q:
            like = f"%{q}%"
            sql += " AND (full_name LIKE ? OR first_name LIKE ? OR last_name LIKE ? OR club LIKE ? OR category_code LIKE ?)"
            args.extend([like, like, like, like, like])
        sql += " ORDER BY COALESCE(rank_position, 999999), full_name"
        with db() as con:
            rows = con.execute(sql, args).fetchall()
        return jsonify([dict(row) for row in rows])

    @app.route("/api/players", methods=["POST"])
    @require_auth
    def create_player():
        payload = clean_player_payload(request.get_json(silent=True) or {})
        with db() as con:
            cur = con.execute("""
                INSERT INTO players(category_code,event_type,gender,age_group,club,rank_position,first_name,last_name,full_name,total_points,tournaments_played,status,notes,created_at,updated_at)
                VALUES(:category_code,:event_type,:gender,:age_group,:club,:rank_position,:first_name,:last_name,:full_name,:total_points,:tournaments_played,:status,:notes,:created_at,:updated_at)
            """, payload)
            con.commit()
            player = con.execute("SELECT * FROM players WHERE id=?", (cur.lastrowid,)).fetchone()
        return jsonify(dict(player)), 201

    @app.route("/api/players/<int:player_id>", methods=["PUT"])
    @require_auth
    def update_player(player_id):
        payload = clean_player_payload(request.get_json(silent=True) or {}, updating=True)
        payload["id"] = player_id
        with db() as con:
            found = con.execute("SELECT id FROM players WHERE id=?", (player_id,)).fetchone()
            if not found:
                return jsonify({"error": "Player not found"}), 404
            con.execute("""
                UPDATE players SET
                    category_code=:category_code,event_type=:event_type,gender=:gender,age_group=:age_group,
                    club=:club,rank_position=:rank_position,first_name=:first_name,last_name=:last_name,
                    full_name=:full_name,total_points=:total_points,tournaments_played=:tournaments_played,
                    status=:status,notes=:notes,updated_at=:updated_at
                WHERE id=:id
            """, payload)
            con.commit()
            player = con.execute("SELECT * FROM players WHERE id=?", (player_id,)).fetchone()
        return jsonify(dict(player))


    @app.route("/api/players/<int:player_id>/points", methods=["POST"])
    @require_auth
    def add_player_points(player_id):
        data = request.get_json(silent=True) or {}
        delta = nullable_float(data.get("points_delta"))
        if delta is None:
            return jsonify({"error": "points_delta is required"}), 400
        note = (data.get("notes") or "Latest tournament points added").strip()
        increment_tournament = bool(data.get("increment_tournament", True))
        with db() as con:
            player = con.execute("SELECT * FROM players WHERE id=?", (player_id,)).fetchone()
            if not player:
                return jsonify({"error": "Player not found"}), 404
            old_points = float(player["total_points"] or 0)
            new_points = old_points + float(delta)
            tournaments_played = int(player["tournaments_played"] or 0) + (1 if increment_tournament else 0)
            now = now_iso()
            con.execute("""
                UPDATE players
                SET total_points=?, tournaments_played=?, updated_at=?
                WHERE id=?
            """, (new_points, tournaments_played, now, player_id))
            con.execute("""
                INSERT INTO audit_logs(actor, action, entity, entity_id, details, created_at)
                VALUES(?,?,?,?,?,?)
            """, (ADMIN_USERNAME, "ADD_POINTS", "players", player_id,
                  f"{player['full_name']}: {old_points} + {delta} = {new_points}. {note}", now))
            con.commit()
            updated = con.execute("SELECT * FROM players WHERE id=?", (player_id,)).fetchone()
        return jsonify({"player": dict(updated), "old_points": old_points, "added_points": delta, "new_points": new_points})

    @app.route("/api/players/<int:player_id>", methods=["DELETE"])
    @require_auth
    def delete_player(player_id):
        with db() as con:
            player = con.execute("SELECT * FROM players WHERE id=?", (player_id,)).fetchone()
            if not player:
                return jsonify({"error": "Player not found"}), 404
            con.execute("DELETE FROM players WHERE id=?", (player_id,))
            con.commit()
        return jsonify({"deleted": True, "player": dict(player)})

    @app.route("/api/players/export.csv")
    @require_auth
    def export_players_csv():
        with db() as con:
            rows = con.execute("SELECT * FROM players ORDER BY COALESCE(rank_position,999999), full_name").fetchall()
        output = io.StringIO()
        writer = csv.writer(output)
        # Export every player field so the downloaded register is a complete record.
        headers = [
            ("Player ID", "id"), ("Source Ranking ID", "source_ranking_id"),
            ("Full Name", "full_name"), ("First Name", "first_name"), ("Last Name", "last_name"),
            ("Gender", "gender"), ("Age Group", "age_group"), ("Event Type", "event_type"),
            ("Category", "category_code"), ("Club", "club"), ("Rank Position", "rank_position"),
            ("Total Points", "total_points"), ("Tournaments Played", "tournaments_played"),
            ("Status", "status"), ("Notes", "notes"), ("Created At", "created_at"),
            ("Updated At", "updated_at")
        ]
        writer.writerow([label for label, _ in headers])
        for row in rows:
            writer.writerow([row[key] for _, key in headers])
        return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "attachment; filename=lba_players.csv"})

    @app.route("/api/players/export.xlsx")
    @require_auth
    def export_players_xlsx():
        with db() as con:
            rows = con.execute("SELECT * FROM players ORDER BY COALESCE(rank_position,999999), full_name").fetchall()
        # Keep the Excel register complete: include all fields stored for each player.
        # Human-readable labels keep the workbook understandable to tournament staff.
        header_map = [
            ("Player ID", "id"), ("Source Ranking ID", "source_ranking_id"),
            ("Full Name", "full_name"), ("First Name", "first_name"), ("Last Name", "last_name"),
            ("Gender", "gender"), ("Age Group", "age_group"), ("Event Type", "event_type"),
            ("Category", "category_code"), ("Club", "club"), ("Rank Position", "rank_position"),
            ("Total Points", "total_points"), ("Tournaments Played", "tournaments_played"),
            ("Status", "status"), ("Notes", "notes"), ("Created At", "created_at"),
            ("Updated At", "updated_at")
        ]
        headers = [label for label, _ in header_map]
        keys = [key for _, key in header_map]

        wb = Workbook()
        wb.properties.title = "LBA Official Player Register"
        wb.properties.subject = "Lesotho Badminton Association player register"
        wb.properties.creator = "LBA Admin System"

        ws = wb.active
        ws.title = "Official Player Register"
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        ws.append(headers)

        # Professional register header.
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor="0F766E")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(
                bottom=Side(style="medium", color="0B5F59")
            )
        ws.row_dimensions[1].height = 30

        for row in rows:
            ws.append([row[key] for key in keys])

        # Make the register easy to read and print.
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
        widths = [11, 18, 30, 18, 20, 12, 12, 20, 15, 24, 14, 14, 20, 13, 42, 22, 22]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        for row in ws.iter_rows(min_row=2):
            row[0].alignment = Alignment(horizontal="center")
            row[1].alignment = Alignment(horizontal="center")
            row[5].alignment = Alignment(horizontal="center")
            row[6].alignment = Alignment(horizontal="center")
            row[8].alignment = Alignment(horizontal="center")
            row[10].alignment = Alignment(horizontal="center")
            row[11].number_format = "0.00"
            row[12].alignment = Alignment(horizontal="center")
            row[13].alignment = Alignment(horizontal="center")
            row[14].alignment = Alignment(vertical="top", wrap_text=True)
            row[15].alignment = Alignment(horizontal="center")
            row[16].alignment = Alignment(horizontal="center")

        if ws.max_row >= 2:
            table = Table(displayName="LBAPlayerRegister", ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium4",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            ws.add_table(table)

        ws.auto_filter.ref = ws.dimensions
        ws.print_title_rows = "1:1"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.center.text = "Lesotho Badminton Association • Official Player Register"
        ws.oddFooter.right.text = "Page &P of &N"

        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)
        return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition":"attachment; filename=lba_players.xlsx"})

    @app.route("/api/players/import-template.xlsx")
    @require_auth
    def import_template():
        wb = Workbook()
        ws = wb.active; ws.title = "Players"
        headers = ["full_name","first_name","last_name","gender","age_group","event_type","category_code","club","rank_position","total_points","tournaments_played","status","notes"]
        ws.append(headers)
        ws.append(["Example Player","Example","Player","Men","U15","Men's Singles","MS,U15","LBA Club",1,100,0,"Active",""])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="0F766E")
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(max(max(len(str(c.value or "")) for c in col)+2,12),34)
        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)
        return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition":"attachment; filename=lba_import_template.xlsx"})

    @app.route("/api/import/excel", methods=["POST"])
    @require_auth
    def import_excel():
        if "file" not in request.files:
            return jsonify({"error": "Upload an Excel file using field name 'file'"}), 400
        file = request.files["file"]
        try:
            imported, skipped = import_workbook(file)
        except Exception as exc:
            return jsonify({"error": f"Could not read Excel file: {exc}"}), 400
        return jsonify({"imported": imported, "skipped": skipped})

    @app.route("/api/draws/random", methods=["POST"])
    @require_auth
    def random_draw():
        data = request.get_json(silent=True) or {}
        title = (data.get("title") or "Random Draw").strip()
        category = (data.get("category_code") or "All").strip()
        draw_type = (data.get("draw_type") or "Singles").strip()
        seed_by_rank = bool(data.get("seed_by_rank", False))
        selected_ids = data.get("player_ids") or []
        selected_ids = [nullable_int(pid) for pid in selected_ids]
        selected_ids = [pid for pid in selected_ids if pid is not None]

        with db() as con:
            if selected_ids:
                placeholders = ",".join("?" for _ in selected_ids)
                sql = f"SELECT * FROM players WHERE status='Active' AND id IN ({placeholders})"
                players = [dict(row) for row in con.execute(sql, selected_ids).fetchall()]
                # Preserve the admin's selected order before shuffling/seeding.
                order = {pid: index for index, pid in enumerate(selected_ids)}
                players.sort(key=lambda row: order.get(row["id"], 999999))
            else:
                args = []
                sql = "SELECT * FROM players WHERE status='Active'"
                if category and category != "All":
                    sql += " AND category_code=?"
                    args.append(category)
                players = [dict(row) for row in con.execute(sql, args).fetchall()]

            if len(players) == 0:
                return jsonify({"error": "No active selected players found for this draw"}), 400

            if seed_by_rank:
                players.sort(key=lambda x: (x.get("rank_position") is None, x.get("rank_position") or 999999))
            else:
                import random
                random.shuffle(players)

            fixtures = build_fixtures(players, draw_type)
            now = now_iso()
            cur = con.execute(
                "INSERT INTO draws(title,category_code,draw_type,created_at) VALUES(?,?,?,?)",
                (title, category, draw_type, now),
            )
            draw_id = cur.lastrowid
            for fixture in fixtures:
                con.execute("""
                    INSERT INTO draw_matches(draw_id,match_no,side_a,side_b,side_a_player_ids,side_b_player_ids,status)
                    VALUES(?,?,?,?,?,?,?)
                """, (draw_id, fixture["match_no"], fixture["side_a"], fixture["side_b"], fixture["side_a_player_ids"], fixture["side_b_player_ids"], "Pending"))
            con.execute("""
                INSERT INTO audit_logs(actor, action, entity, entity_id, details, created_at)
                VALUES(?,?,?,?,?,?)
            """, (ADMIN_USERNAME, "GENERATE_DRAW", "draws", draw_id,
                  f"{title}: {len(players)} selected participant(s), {len(fixtures)} match(es)", now))
            con.commit()
            draw = get_draw(con, draw_id)
            draw["selected_player_count"] = len(players)
        return jsonify(draw), 201

    @app.route("/api/draws")
    @require_auth
    def list_draws():
        with db() as con:
            rows = con.execute("SELECT * FROM draws ORDER BY id DESC").fetchall()
        return jsonify([dict(row) for row in rows])

    @app.route("/api/draws/<int:draw_id>")
    @require_auth
    def read_draw(draw_id):
        with db() as con:
            draw = get_draw(con, draw_id)
            if not draw:
                return jsonify({"error": "Draw not found"}), 404
        return jsonify(draw)

    @app.route("/api/draws/<int:draw_id>/export.pdf")
    @require_auth
    def export_draw_pdf(draw_id):
        with db() as con:
            draw = get_draw(con, draw_id)
            if not draw:
                return jsonify({"error": "Draw not found"}), 404

        pdf_bytes = build_draw_pdf(draw)
        filename = safe_filename(f"{draw['title'] or 'lba_draw'}_{draw['id']}.pdf")
        return Response(
            pdf_bytes,
            mimetype="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    @app.route("/api/draws/<int:draw_id>", methods=["DELETE"])
    @require_auth
    def delete_draw(draw_id):
        with db() as con:
            con.execute("DELETE FROM draw_matches WHERE draw_id=?", (draw_id,))
            con.execute("DELETE FROM draws WHERE id=?", (draw_id,))
            con.commit()
        return jsonify({"deleted": True})

    return app


def send_email(to_address, subject, body):
    # Railway deployments may block outbound SMTP connections. Prefer an HTTPS
    # email API when configured, while retaining SMTP as a fallback for local use.
    if RESEND_API_KEY:
        payload = json.dumps({
            "from": EMAIL_FROM,
            "to": [to_address],
            "subject": subject,
            "text": body,
        }).encode("utf-8")
        req = urllib.request.Request(
            "https://api.resend.com/emails",
            data=payload,
            headers={
                "Authorization": f"Bearer {RESEND_API_KEY}",
                "Content-Type": "application/json",
                "User-Agent": "LBA-Mobile-App/1.0",
            },
            method="POST",
        )
        print(f"[EMAIL] Resend email attempt: to={to_address!r}, subject={subject!r}", flush=True)
        try:
            with urllib.request.urlopen(req, timeout=20) as response:
                detail = response.read().decode("utf-8", errors="replace")
                print(f"[EMAIL] Resend response status: {response.status}", flush=True)
                print(f"[EMAIL] Resend response body: {detail[:1000]}", flush=True)
                if response.status < 200 or response.status >= 300:
                    raise RuntimeError(f"Email API returned HTTP {response.status}")
                return
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            print(f"[EMAIL] Resend HTTP error: status={exc.code}, body={detail[:1000]}", flush=True)
            raise RuntimeError(f"Email API returned HTTP {exc.code}: {detail[:500]}") from exc
        except urllib.error.URLError as exc:
            print(f"[EMAIL] Resend email exception: {exc.reason!r}", flush=True)
            raise RuntimeError(f"Could not reach email API: {exc.reason}") from exc
        except Exception as exc:
            print(f"[EMAIL] Resend email exception: {type(exc).__name__}: {exc}", flush=True)
            raise

    if not SMTP_HOST or not SMTP_USERNAME or not SMTP_PASSWORD or not SMTP_FROM:
        raise RuntimeError(
            "Email service is not configured. Set RESEND_API_KEY and LBA_EMAIL_FROM, "
            "or configure the LBA_SMTP_* variables."
        )

    message = EmailMessage()
    message["From"] = SMTP_FROM
    message["To"] = to_address
    message["Subject"] = subject
    message.set_content(body)

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=20) as server:
        server.starttls()
        server.login(SMTP_USERNAME, SMTP_PASSWORD)
        server.send_message(message)


def bearer_token():
    auth = request.headers.get("Authorization", "")
    return auth.replace("Bearer ", "", 1).strip()


def public_user(row):
    if not row:
        return None
    return {
        "id": row["id"],
        "username": row["username"],
        "email": row["email"],
        "full_name": row["full_name"],
        "role": row["role"],
        "is_active": bool(row["is_active"]),
        "email_verified": bool(row["email_verified"]),
        "created_at": row["created_at"],
        "last_login": row["last_login"],
    }


def require_auth(func):
    @wraps(func)
    def wrapper(*args, **kwargs):
        token = bearer_token()
        if not token:
            return jsonify({"error": "Unauthorized"}), 401

        token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
        with db() as con:
            user = con.execute("""
                SELECT u.*
                FROM auth_tokens t
                JOIN users u ON u.id = t.user_id
                WHERE t.token_hash=?
                  AND t.expires_at > ?
                  AND u.is_active=1
            """, (token_hash, now_iso())).fetchone()

        if not user:
            return jsonify({"error": "Unauthorized"}), 401

        g.current_user = user
        return func(*args, **kwargs)
    return wrapper


def db():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con


def now_iso():
    return datetime.utcnow().replace(microsecond=0).isoformat() + "Z"


def ensure_database():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with db() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS players (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_ranking_id INTEGER,
            category_code TEXT,
            event_type TEXT,
            gender TEXT,
            age_group TEXT,
            club TEXT,
            rank_position INTEGER,
            first_name TEXT,
            last_name TEXT,
            full_name TEXT NOT NULL,
            total_points REAL DEFAULT 0,
            tournaments_played INTEGER DEFAULT 0,
            status TEXT DEFAULT 'Active',
            notes TEXT,
            created_at TEXT,
            updated_at TEXT
        );
        CREATE TABLE IF NOT EXISTS draws (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            category_code TEXT,
            draw_type TEXT,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS draw_matches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            draw_id INTEGER NOT NULL,
            match_no INTEGER NOT NULL,
            side_a TEXT NOT NULL,
            side_b TEXT NOT NULL,
            side_a_player_ids TEXT,
            side_b_player_ids TEXT,
            winner TEXT,
            status TEXT DEFAULT 'Pending',
            FOREIGN KEY(draw_id) REFERENCES draws(id)
        );
        CREATE TABLE IF NOT EXISTS audit_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            actor TEXT,
            action TEXT,
            entity TEXT,
            entity_id INTEGER,
            details TEXT,
            created_at TEXT
        );
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            email TEXT NOT NULL UNIQUE,
            full_name TEXT NOT NULL,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'Viewer',
            is_active INTEGER NOT NULL DEFAULT 1,
            email_verified INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            last_login TEXT
        );
        CREATE TABLE IF NOT EXISTS auth_tokens (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            token_hash TEXT NOT NULL UNIQUE,
            created_at TEXT NOT NULL,
            expires_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
        );
        CREATE INDEX IF NOT EXISTS idx_auth_tokens_hash ON auth_tokens(token_hash);
        CREATE INDEX IF NOT EXISTS idx_auth_tokens_user ON auth_tokens(user_id);
        CREATE TABLE IF NOT EXISTS email_tokens (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            token_hash TEXT NOT NULL UNIQUE,
            token_type TEXT NOT NULL,
            expires_at TEXT NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
        );
        CREATE INDEX IF NOT EXISTS idx_email_tokens_hash ON email_tokens(token_hash);
        """)
        admin = con.execute("SELECT id FROM users WHERE username=?", (ADMIN_USERNAME,)).fetchone()
        if not admin and ADMIN_PASSWORD:
            now = now_iso()
            con.execute("""
                INSERT INTO users(username,email,full_name,password_hash,role,is_active,email_verified,created_at,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?)
            """, (
                ADMIN_USERNAME,
                os.environ.get("LBA_ADMIN_EMAIL", "admin@lba.local"),
                "LBA Administrator",
                generate_password_hash(ADMIN_PASSWORD),
                "Super Admin",
                1,
                1,
                now,
                now,
            ))

        count = con.execute("SELECT COUNT(*) FROM players").fetchone()[0]
        has_combined = con.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='combined_rankings'").fetchone()
        if count == 0 and has_combined:
            now = now_iso()
            con.execute("""
                INSERT INTO players(source_ranking_id,category_code,event_type,gender,age_group,club,rank_position,first_name,last_name,full_name,total_points,tournaments_played,status,created_at,updated_at)
                SELECT id, category_code, event_type, gender, age_group, COALESCE(club,''), rank_position, first_name, last_name, full_name,
                       COALESCE(total_points,0), COALESCE(tournaments_played,0), 'Active', ?, ?
                FROM combined_rankings
            """, (now, now))
        con.commit()


def clean_player_payload(data, updating=False):
    full_name = (data.get("full_name") or data.get("name") or "").strip()
    first_name = (data.get("first_name") or "").strip()
    last_name = (data.get("last_name") or "").strip()
    if not full_name:
        full_name = f"{first_name} {last_name}".strip()
    if not full_name:
        raise ValueError("full_name is required")
    if not first_name and full_name:
        parts = full_name.split()
        first_name = parts[0]
        last_name = " ".join(parts[1:]) if len(parts) > 1 else last_name
    category_code = (data.get("category_code") or data.get("category") or "Uncategorised").strip()
    event_type = (data.get("event_type") or infer_event_type(category_code)).strip()
    gender = (data.get("gender") or infer_gender(category_code)).strip()
    age_group = (data.get("age_group") or infer_age_group(category_code)).strip()
    now = now_iso()
    return {
        "category_code": category_code,
        "event_type": event_type,
        "gender": gender,
        "age_group": age_group,
        "club": (data.get("club") or "").strip(),
        "rank_position": nullable_int(data.get("rank_position", data.get("rank"))),
        "first_name": first_name,
        "last_name": last_name,
        "full_name": full_name,
        "total_points": nullable_float(data.get("total_points", data.get("points"))) or 0,
        "tournaments_played": nullable_int(data.get("tournaments_played")) or 0,
        "status": data.get("status") or "Active",
        "notes": data.get("notes") or "",
        "created_at": now,
        "updated_at": now,
    }


def nullable_int(value):
    if value in (None, ""):
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def nullable_float(value):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (ValueError, TypeError):
        return None


def infer_gender(category):
    text = (category or "").upper()
    if text.startswith("WS") or "WOM" in text:
        return "Women"
    if text.startswith("MS") or "MEN" in text:
        return "Men"
    return "Open"


def infer_event_type(category):
    text = (category or "").upper()
    if text.startswith("WS"):
        return "Women's Singles"
    if text.startswith("MS"):
        return "Men's Singles"
    if "XD" in text or "MIX" in text:
        return "Mixed Doubles"
    if "D" in text:
        return "Doubles"
    return "Singles"


def infer_age_group(category):
    text = (category or "").upper().replace(" ", "")
    for age in ["U11", "U13", "U15", "U17", "U19", "18+"]:
        if age in text:
            return age
    return "Open"


def build_fixtures(players, draw_type):
    fixtures = []
    if draw_type.lower() == "doubles":
        pairs = []
        for i in range(0, len(players), 2):
            a = players[i]
            b = players[i + 1] if i + 1 < len(players) else None
            name = f"{a['full_name']} / {b['full_name']}" if b else f"{a['full_name']} / Waiting partner"
            ids = f"{a['id']},{b['id']}" if b else f"{a['id']}"
            pairs.append({"name": name, "ids": ids})
        for i in range(0, len(pairs), 2):
            a = pairs[i]
            b = pairs[i + 1] if i + 1 < len(pairs) else {"name": "Bye", "ids": ""}
            fixtures.append({"match_no": len(fixtures) + 1, "side_a": a["name"], "side_b": b["name"], "side_a_player_ids": a["ids"], "side_b_player_ids": b["ids"]})
        return fixtures

    for i in range(0, len(players), 2):
        a = players[i]
        b = players[i + 1] if i + 1 < len(players) else {"full_name": "Bye", "id": ""}
        fixtures.append({"match_no": len(fixtures) + 1, "side_a": a["full_name"], "side_b": b["full_name"], "side_a_player_ids": str(a["id"]), "side_b_player_ids": str(b["id"])})
    return fixtures


def get_draw(con, draw_id):
    row = con.execute("SELECT * FROM draws WHERE id=?", (draw_id,)).fetchone()
    if not row:
        return None
    matches = con.execute("SELECT * FROM draw_matches WHERE draw_id=? ORDER BY match_no", (draw_id,)).fetchall()
    data = dict(row)
    data["matches"] = [dict(match) for match in matches]
    return data


def safe_filename(name):
    cleaned = re.sub(r"[^A-Za-z0-9_.-]+", "_", name).strip("_")
    return cleaned or "lba_draw.pdf"


def build_draw_pdf(draw):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=draw.get("title") or "LBA Draw",
    )
    styles = getSampleStyleSheet()
    story = []

    title_style = styles["Title"]
    title_style.fontName = "Helvetica-Bold"
    title_style.fontSize = 18
    title_style.leading = 22
    title_style.textColor = colors.HexColor("#0f172a")

    normal = styles["BodyText"]
    normal.fontName = "Helvetica"
    normal.fontSize = 9
    normal.leading = 12

    small = styles["BodyText"]
    small.fontName = "Helvetica"
    small.fontSize = 8
    small.leading = 10
    small.textColor = colors.HexColor("#475569")

    heading_style = styles["Heading2"]
    heading_style.textColor = colors.HexColor("#047857")
    heading_style.spaceAfter = 4

    logo_path = BASE_DIR.parent / "frontend" / "src" / "assets" / "lba-logo.png"
    if logo_path.exists():
        try:
            story.append(Image(str(logo_path), width=26 * mm, height=26 * mm))
            story.append(Spacer(1, 3 * mm))
        except Exception:
            pass

    story.append(Paragraph("Lesotho Badminton Association", title_style))
    story.append(Paragraph("Official Tournament Draw Sheet", heading_style))
    story.append(Spacer(1, 5 * mm))

    meta = [
        ["Draw title", draw.get("title") or "-", "Draw type", draw.get("draw_type") or "-"],
        ["Category", draw.get("category_code") or "All", "Created", draw.get("created_at") or "-"],
        ["Draw ID", str(draw.get("id") or "-"), "Matches", str(len(draw.get("matches") or []))],
    ]
    meta_table = Table(meta, colWidths=[25 * mm, 65 * mm, 25 * mm, 65 * mm])
    meta_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#cbd5e1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f0")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 7 * mm))

    # Present the draw in normal sports language rather than database terminology.
    # A doubles fixture already contains both players on each side, e.g.
    # "Player A / Player B vs Player C / Player D".
    match_rows = [["Match", "Players", "Winner / Score"]]
    for match in draw.get("matches") or []:
        player_a = str(match.get("side_a") or "-")
        player_b = str(match.get("side_b") or "-")
        fixture = f"{player_a}  vs.  {player_b}"
        match_rows.append([
            str(match.get("match_no") or ""),
            Paragraph(fixture, normal),
            "",
        ])

    table = Table(match_rows, colWidths=[18 * mm, 124 * mm, 40 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#cbd5e1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.append(table)
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("Prepared by LBA Admin System. Keep this file as the official draw record for the tournament.", small))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def import_workbook(file_storage):
    wb = load_workbook(file_storage, data_only=True)
    imported = 0
    skipped = 0
    now = now_iso()

    with db() as con:
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            # LBA ranking workbooks may have title/metadata rows before the
            # actual player-table header. Find the row containing Last Name
            # and First Name instead of assuming the first row is the header.
            header_index = None
            for idx, row in enumerate(rows):
                normalized = [
                    str(cell).strip().lower().replace(" ", "_") if cell is not None else ""
                    for cell in row
                ]
                if "last_name" in normalized and "first_name" in normalized:
                    header_index = idx
                    break

            if header_index is None:
                # Also support simple import sheets whose first row is already
                # a standard header.
                header_index = 0

            headers = [
                str(cell).strip().lower().replace(" ", "_") if cell is not None else ""
                for cell in rows[header_index]
            ]

            # The official LBA ranking workbook stores:
            # column B = rank, C = last name, D = first name,
            # E = total points, F = tournaments played.
            def get_value(record, *names):
                for name in names:
                    value = record.get(name)
                    if value not in (None, ""):
                        return value
                return None

            # Infer category/event information from the worksheet name.
            sheet_name = str(sheet.title).strip()
            category_from_sheet = sheet_name.replace(" ", "")
            if category_from_sheet.upper().startswith("MS,U"):
                category_from_sheet = category_from_sheet.upper()
            elif category_from_sheet.upper().startswith("MS,18"):
                category_from_sheet = "MS,18+"
            elif category_from_sheet.upper().startswith("WS,U"):
                category_from_sheet = category_from_sheet.upper()
            elif category_from_sheet.upper().startswith("WS,18"):
                category_from_sheet = "WS,18+"

            for values in rows[header_index + 1:]:
                record = dict(zip(headers, values))

                full_name = str(get_value(
                    record, "full_name", "name", "player_name",
                    "player", "player_full_name"
                ) or "").strip()

                first = str(get_value(
                    record, "first_name", "firstname", "first"
                ) or "").strip()
                last = str(get_value(
                    record, "last_name", "lastname", "surname", "last"
                ) or "").strip()

                # Official LBA workbook uses C/D for last/first name.
                if not full_name:
                    full_name = f"{first} {last}".strip()

                if not full_name:
                    skipped += 1
                    continue

                rank = get_value(
                    record, "rank_position", "rank", "#"
                )
                total_points = get_value(
                    record, "total_points", "points", "points_total"
                )
                tournaments_played = get_value(
                    record, "tournaments_played", "tournaments", "tournament_played"
                )

                category = str(get_value(
                    record, "category_code", "category", "category_name",
                    "event_category"
                ) or category_from_sheet).strip()

                payload = clean_player_payload({
                    "full_name": full_name,
                    "first_name": first,
                    "last_name": last,
                    "category_code": category,
                    "gender": get_value(record, "gender"),
                    "age_group": get_value(record, "age_group"),
                    "event_type": get_value(record, "event_type"),
                    "club": get_value(record, "club"),
                    "rank_position": rank,
                    "total_points": total_points,
                    "tournaments_played": tournaments_played,
                    "status": "Active",
                    "notes": f"Imported from {sheet.title}",
                })

                con.execute("""
                    INSERT INTO players(
                        category_code,event_type,gender,age_group,club,
                        rank_position,first_name,last_name,full_name,
                        total_points,tournaments_played,status,notes,
                        created_at,updated_at
                    )
                    VALUES(
                        :category_code,:event_type,:gender,:age_group,:club,
                        :rank_position,:first_name,:last_name,:full_name,
                        :total_points,:tournaments_played,:status,:notes,
                        :created_at,:updated_at
                    )
                """, payload)
                imported += 1

        con.commit()

    return imported, skipped


app = create_app()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5050"))
    app.run(debug=False, host="0.0.0.0", port=port)
